"""Generate disposable synthetic fixtures, containing no personal knowledge."""
import io
import sys
from pathlib import Path
from docx import Document
from docx.oxml import parse_xml
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, NumberObject, DecodedStreamObject

root = Path(sys.argv[1])
root.mkdir(parents=True, exist_ok=True)
image = io.BytesIO()
Image.new('RGB', (12, 8), (30, 100, 180)).save(image, format='PNG')
(root / 'expected.png').write_bytes(image.getvalue())
doc = Document()
doc.add_heading('中文格式验收', 1)
doc.add_paragraph('第一段：<script>不是可执行内容</script>，$& 原样保留。')
table = doc.add_table(rows=2, cols=2)
for cell, value in zip([cell for row in table.rows for cell in row.cells], ['项目', '数量', '样本', '2']):
    cell.text = value
doc.add_picture(io.BytesIO(image.getvalue()))
doc.add_paragraph('末段：表格与图片之后。')
box_paragraph = doc.add_paragraph()
box_paragraph._p.append(parse_xml('''<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t xml:space="preserve">    if (value &lt; 9) {</w:t></w:r></w:p><w:p><w:r><w:t xml:space="preserve">        System.gc();</w:t><w:br/><w:t>    }</w:t></w:r></w:p><w:p><w:r><w:t>&lt;script&gt;literal&lt;/script&gt;</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>'''))
doc.save(root / 'sample.docx')

writer = PdfWriter()
font = writer._add_object(DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')}))
picture = DecodedStreamObject()
picture.set_data(bytes([30,100,180])*12*8)
picture.update({NameObject('/Type'):NameObject('/XObject'),NameObject('/Subtype'):NameObject('/Image'),NameObject('/Width'):NumberObject(12),NameObject('/Height'):NumberObject(8),NameObject('/ColorSpace'):NameObject('/DeviceRGB'),NameObject('/BitsPerComponent'):NumberObject(8)})
picture_ref = writer._add_object(picture)
for content in [
    b'BT /F1 18 Tf 40 730 Td (TEXT PAGE TITLE) Tj 0 -35 Td (First paragraph before table.) Tj 0 -35 Td (Name       Value) Tj 0 -25 Td (Sample     2) Tj ET',
    b'q 240 0 0 160 30 500 cm /Image Do Q',
    b'',
    b'BT /F1 18 Tf 40 730 Td (MIXED PAGE) Tj ET q 120 0 0 80 30 500 cm /Image Do Q'
]:
    page = writer.add_blank_page(width=595,height=842)
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font}),NameObject('/XObject'):DictionaryObject({NameObject('/Image'):picture_ref} if b'/Image Do' in content else {})})
    stream = DecodedStreamObject(); stream.set_data(content)
    page[NameObject('/Contents')] = writer._add_object(stream)
writer.write(root / 'sample.pdf')
bad_page = writer.add_blank_page(width=595,height=842)
bad_stream = DecodedStreamObject(); bad_stream.set_data(b'partial fixture')
bad_stream[NameObject('/Filter')] = NameObject('/UnsupportedFixtureFilter')
bad_page[NameObject('/Contents')] = writer._add_object(bad_stream)
writer.write(root / 'partial.pdf')
writer.encrypt('fixture-password')
writer.write(root / 'encrypted.pdf')
